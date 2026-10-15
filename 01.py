# -*- coding: utf-8 -*-
"""
Скрипт форматирования Excel-документов для электронного архива (openpyxl + COM).

Логика:
1. Распаковка архивов (.zip встроенным zipfile, .rar/.7z — внешним распаковщиком).
2. openpyxl — очистка заливок, поиск шапки, сквозные строки, колонтитул,
   РАСЧЁТ границ области печати, ЗАПОМИНАНИЕ исходных настроек страницы.
3. COM (Excel) — установка области печати и сквозных строк. Дальше:
   - если оригинал УЖЕ помещается на A4 → восстанавливаем его поля,
     ориентацию и масштаб и больше ничего не трогаем;
   - иначе — перенос текста, подгонка высоты шапки, подбор ориентации
     и масштаба по реальной ширине от Excel.

ПРАВИЛО «НЕ ЧИНИ ИСПРАВНОЕ» (PRESERVE_IF_FITS):
если с родными полями, ориентацией и масштабом таблица влезает по ширине,
файл оставляем как есть. Тот, кто его делал, уже всё настроил.
В этой проверке PRINTER_MARGIN_PTS НЕ вычитается намеренно: файл уже
существует и печатался в таком виде, поправка на принтер в нём учтена
автором. Вместо неё — небольшой допуск PRESERVE_TOLERANCE на погрешность
измерения ширины.
Не сохраняем, если: у листа включено «вписать в страницу» (тогда масштаб
в файле фиктивный) или исходный масштаб ниже MIN_SCALE.

ПРАВИЛО ОРИЕНТАЦИИ ПО ФОРМЕ ТАБЛИЦЫ (когда переделываем):
  - вытянута в ширину  → альбомная. Если не помещается на страницу даже
    в альбомной — масштаб под страницу НЕ подбираем, лист печатается
    лентой на несколько страниц;
  - вытянута в высоту  → книжная.
Ориентацию решает форма таблицы, а не то, что стояло в файле. Прежняя
логика «переворачиваем, если не влезает» была неверна: она переворачивала
альбомную в книжную, где страница уже и масштаб заведомо мельче.

ПОЧЕМУ ОБЛАСТЬ ПЕЧАТИ СТАВИТ COM, А НЕ OPENPYXL:
некоторые версии openpyxl записывают имя как "Print_Area" вместо
"_xlnm.Print_Area". Excel такое имя игнорирует, область печати не работает,
и ширина меряется по фантомным колонкам (после конвертации .xls → .xlsx лист
часто тянется на десятки пустых колонок).

ПОЧЕМУ ВЫСОТУ СТРОК ШАПКИ СЧИТАЕМ САМИ:
Rows.AutoFit() не учитывает объединённые ячейки — это ограничение Excel.
Включённый перенос по словам меняет разбивку текста, высота при этом
не растёт, и лишние строки уходят под обрез. Нужную высоту измеряем
в служебной книге. Только увеличиваем — вручную подогнанные не портим.

ВАЖНО про RAR: модуль rarfile — только обёртка, ему нужен внешний
распаковщик. Скрипт ищет UnRAR.exe (WinRAR) или 7z.exe (7-Zip) сам.
"""
from pathlib import Path
import zipfile
import logging
import os
import shutil
import subprocess
import sys
import time
import tkinter as tk
import traceback
from tkinter import filedialog, messagebox, ttk
import re

# Excel / COM
import pythoncom
import win32com.client as win32

# openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.styles import PatternFill, Border

# Настройка логгера
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger(__name__)

# -----------------------
# Константы
# -----------------------
PAGE_WIDTH_A4_PORTRAIT = 8.27       # A4, дюймы
PAGE_HEIGHT_A4_PORTRAIT = 11.69     # A4, дюймы
MAX_FILES_TO_PROCESS = 1000

# Ориентация и масштаб
XL_PORTRAIT, XL_LANDSCAPE = 1, 2    # xlPortrait / xlLandscape
XL_PAPER_A4 = 9                     # xlPaperA4
MIN_SCALE = 60                      # ниже не опускаемся
# Порог используется ТОЛЬКО когда исходная ориентация неизвестна
PORTRAIT_PREFERENCE = 60

# --- Правило ориентации по форме таблицы ---
# Вытянутой в высоту считаем таблицу, у которой высота области печати
# больше ширины во столько раз. 1.0 = просто выше, чем шире.
TALL_TABLE_RATIO = 1.0

# Ориентация для таблиц, вытянутых в высоту.
# None -> XL_PORTRAIT. Поставьте XL_LANDSCAPE, если нужна альбомная.
TALL_TABLE_ORIENTATION = None

# Масштаб широкой таблицы, которая не влезает на страницу:
#   True  — вернуть авторский масштаб из оригинала (если он вменяемый)
#   False — оставить 100 %
WIDE_TABLE_KEEP_ORIGINAL_SCALE = True

# --- «Не чини исправное» ---
# Если оригинал уже помещается на A4 — не менять поля, ориентацию, масштаб.
PRESERVE_IF_FITS = True
# Допуск на погрешность измерения ширины (1.02 = 2%).
# Больше значение → чаще оставляем файл как есть.
PRESERVE_TOLERANCE = 1.02
# Включать перенос текста и подгонку шапки в сохранённых файлах.
# По умолчанию False: раз оригинал печатался нормально, не вмешиваемся.
WRAP_TEXT_WHEN_PRESERVED = False

# Константы для Cells.Find (резервный способ определения данных)
XL_FORMULAS = -4123                 # xlFormulas
XL_PREVIOUS = 2                     # xlPrevious
XL_BY_ROWS = 1                      # xlByRows
XL_BY_COLUMNS = 2                   # xlByColumns

# Непечатаемая зона принтера, pts, суммарно слева+справа.
# Применяется ТОЛЬКО когда пересчитываем масштаб сами.
#
# КАЛИБРОВКА: поставьте 0, обработайте пару файлов, откройте результат
# и посмотрите в строке состояния Excel "Страница: 1 из N".
#   - если по ширине лист режется на две страницы → увеличивайте на 10–15
#   - если всё влезает, а масштаб хочется крупнее → уменьшайте
# 60 pts ≈ 0.42" с каждой стороны — для драйвера 1200 dpi.
# 30 pts ≈ 0.21" с каждой стороны — для 600 dpi.
PRINTER_MARGIN_PTS = 60.0

# Наши поля страницы (дюймы) — применяются только при переформатировании
MARGIN_LEFT = 0.25
MARGIN_RIGHT = 0.25
MARGIN_TOP = 0.5
MARGIN_BOTTOM = 0.75
MARGIN_HEADER = 0.25
MARGIN_FOOTER = 0.25

# Запас строк снизу в области печати. По колонкам запаса быть НЕ должно:
# каждая лишняя пустая колонка напрямую уменьшает масштаб печати.
PRINT_AREA_ROW_PADDING = 2
# Потолок сканирования строк. Не режет данные — от бесконечных листов
# защищает ранняя остановка после 5 пустых строк подряд.
MAX_SCAN_ROWS = 50000
MAX_SCAN_COLS = 100

# Включать перенос по словам в области печати (при переформатировании)
ENABLE_WRAP_TEXT = True

# Подгонять высоту строк шапки под перенесённый текст
FIX_HEADER_HEIGHTS = True
HEIGHT_TOLERANCE_PTS = 1.5
MAX_HEADER_ROW_HEIGHT = 220.0

# Поддерживаемые расширения архивов
ARCHIVE_EXTENSIONS = (".zip", ".rar", ".7z")

# Скрыть окно консоли при вызове распаковщика (только Windows)
_NO_WINDOW = 0x08000000 if sys.platform == "win32" else 0


# -----------------------
# Поиск внешнего распаковщика
# -----------------------
_ARCHIVE_TOOL_CACHE = "not_searched"


def find_archive_tool():
    """Ищет UnRAR.exe (WinRAR) или 7z.exe (7-Zip). Результат кэшируется."""
    global _ARCHIVE_TOOL_CACHE
    if _ARCHIVE_TOOL_CACHE != "not_searched":
        return _ARCHIVE_TOOL_CACHE

    candidates = [
        shutil.which("unrar"),
        shutil.which("UnRAR"),
        r"C:\Program Files\WinRAR\UnRAR.exe",
        r"C:\Program Files (x86)\WinRAR\UnRAR.exe",
        r"C:\Program Files\WinRAR\Rar.exe",
        shutil.which("7z"),
        r"C:\Program Files\7-Zip\7z.exe",
        r"C:\Program Files (x86)\7-Zip\7z.exe",
        r"C:\Program Files\7-Zip\7za.exe",
    ]
    for candidate in candidates:
        if candidate and os.path.exists(candidate):
            _ARCHIVE_TOOL_CACHE = candidate
            logger.info(f"Распаковщик архивов: {candidate}")
            return candidate

    _ARCHIVE_TOOL_CACHE = None
    logger.warning(
        "Не найден распаковщик для RAR/7z. Установите WinRAR или 7-Zip, "
        "либо пропишите путь к UnRAR.exe / 7z.exe в find_archive_tool()."
    )
    return None


def extract_with_external_tool(archive_path, extract_to):
    """Распаковывает .rar/.7z внешним распаковщиком"""
    tool = find_archive_tool()
    if tool is None:
        logger.error(f"Нечем распаковать {archive_path.name} — распаковщик не найден")
        return False

    tool_name = os.path.basename(tool).lower()
    if tool_name.startswith(("unrar", "rar")):
        cmd = [tool, "x", "-y", "-idq", str(archive_path), str(extract_to) + os.sep]
    else:
        # 7-Zip: -o без пробела перед путём — так требует синтаксис
        cmd = [tool, "x", "-y", f"-o{extract_to}", str(archive_path)]

    try:
        result = subprocess.run(
            cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
            creationflags=_NO_WINDOW, timeout=600,
        )
        if result.returncode == 0:
            return True
        output = (result.stdout or b"").decode("cp866", errors="replace").strip()
        logger.error(f"Распаковщик вернул код {result.returncode} для "
                     f"{archive_path.name}: {output[:300]}")
        return False
    except subprocess.TimeoutExpired:
        logger.error(f"Распаковка {archive_path.name} прервана по таймауту")
        return False
    except Exception as e:
        logger.error(f"Ошибка вызова распаковщика для {archive_path.name}: {e}")
        return False


def extract_zip_with_correct_encoding(archive_path, extract_to):
    """Распаковывает .zip с восстановлением русских имён файлов"""
    encodings_to_try = ['cp866', 'cp1251']
    extracted = []
    try:
        with zipfile.ZipFile(archive_path, 'r') as zip_ref:
            for info in zip_ref.infolist():
                name = info.filename
                # бит 0x800 означает, что имя уже в UTF-8 — не перекодируем
                if not (info.flag_bits & 0x800):
                    raw = name.encode('cp437', errors='replace')
                    for encoding in encodings_to_try:
                        try:
                            decoded = raw.decode(encoding)
                            if decoded.strip():
                                name = decoded
                                break
                        except Exception:
                            continue

                target_path = extract_to / name
                if info.is_dir():
                    target_path.mkdir(parents=True, exist_ok=True)
                    continue
                target_path.parent.mkdir(parents=True, exist_ok=True)
                with zip_ref.open(info) as source, open(target_path, 'wb') as target:
                    target.write(source.read())
                extracted.append(target_path)
        logger.info(f"ZIP распакован ({len(extracted)} файлов)")
        return True
    except Exception as e:
        logger.error(f"Ошибка распаковки ZIP {archive_path.name}: {e}")
        return False


def extract_archive(archive_path, output_folder=None, _depth=0):
    """Распаковывает архив в папку с его именем, рекурсивно обрабатывая вложенные"""
    archive_path = Path(archive_path)
    if output_folder is None:
        output_folder = archive_path.parent
    extract_to = Path(output_folder) / archive_path.stem
    suffix = archive_path.suffix.lower()

    if _depth > 5:
        logger.warning(f"Слишком глубокая вложенность архивов: {archive_path.name}")
        return None

    try:
        extract_to.mkdir(parents=True, exist_ok=True)
        logger.info(f"Распаковка архива: {archive_path.name} → {extract_to.name}/")

        if suffix == ".zip":
            ok = extract_zip_with_correct_encoding(archive_path, extract_to)
        elif suffix in (".rar", ".7z"):
            ok = extract_with_external_tool(archive_path, extract_to)
            if ok:
                count = sum(1 for _ in extract_to.rglob("*") if _.is_file())
                logger.info(f"{suffix.upper()[1:]} распакован ({count} файлов)")
        else:
            logger.warning(f"Формат {suffix} не поддерживается")
            ok = False

        if not ok:
            try:
                if not any(extract_to.iterdir()):
                    extract_to.rmdir()
            except Exception:
                pass
            return None

        for item in list(extract_to.rglob("*")):
            if item.is_file() and item.suffix.lower() in ARCHIVE_EXTENSIONS:
                logger.info(f"Найден вложенный архив: {item.name}")
                extract_archive(item, item.parent, _depth + 1)

        return extract_to
    except Exception as e:
        logger.error(f"Ошибка распаковки архива {archive_path.name}: {e}")
        return None


def process_archives_in_folder(folder_path, delete_originals=False):
    """Находит и распаковывает все архивы в папке и подпапках"""
    folder_path = Path(folder_path)
    processed_folders = []

    root_archives = [item for item in folder_path.iterdir()
                     if item.is_file() and item.suffix.lower() in ARCHIVE_EXTENSIONS]
    logger.info(f"Найдено архивов в '{folder_path.name}': {len(root_archives)}")

    for archive in root_archives:
        try:
            extracted_folder = extract_archive(archive)
            if extracted_folder:
                processed_folders.append(extracted_folder)
                if delete_originals:
                    try:
                        archive.unlink()
                        logger.info(f"Удален исходный архив: {archive.name}")
                    except Exception as e:
                        logger.warning(f"Не удалось удалить архив {archive.name}: {e}")
                else:
                    logger.info(f"Оригинальный архив сохранен: {archive.name}")
            else:
                logger.warning(f"Архив {archive.name} распаковать не удалось")
        except Exception as e:
            logger.error(f"Ошибка обработки архива {archive.name}: {e}")
            continue

    for subfolder in folder_path.iterdir():
        if subfolder.is_dir():
            if any(subfolder.name == f.stem for f in root_archives):
                continue
            processed_folders.extend(process_archives_in_folder(subfolder, delete_originals))

    return processed_folders


def clean_empty_folders(folder_path):
    try:
        deleted_count = 0
        for item in folder_path.iterdir():
            if item.is_dir():
                deleted_count += clean_empty_folders(item)
                try:
                    if not any(item.iterdir()):
                        item.rmdir()
                        deleted_count += 1
                except Exception:
                    pass
        return deleted_count
    except Exception as e:
        logger.warning(f"Ошибка при очистке пустых папок: {e}")
        return 0


# -----------------------
# Вспомогательные для Excel (openpyxl)
# -----------------------
PRINT_NAMES_TO_CLEAN = (
    '_xlnm.Print_Area', '_xlnm.Print_Titles',
    'Print_Area', 'Print_Titles',
)


def cleanup_old_print_settings(wb):
    """Удаляет старые определения печати, включая имена без префикса _xlnm."""
    removed = 0
    try:
        try:
            dn = wb.defined_names
            if hasattr(dn, 'definedName'):                    # openpyxl < 3.1
                for name in list(dn.definedName):
                    if getattr(name, 'name', None) in PRINT_NAMES_TO_CLEAN:
                        try:
                            dn.definedName.remove(name)
                            removed += 1
                        except Exception:
                            pass
            else:                                             # openpyxl >= 3.1
                for key in list(dn.keys()):
                    if key in PRINT_NAMES_TO_CLEAN:
                        try:
                            del dn[key]
                            removed += 1
                        except Exception:
                            pass
        except Exception as e:
            logger.debug(f"Чистка имён книги: {e}")

        for ws in wb.worksheets:
            try:
                wsdn = getattr(ws, 'defined_names', None)
                if wsdn is None:
                    continue
                for key in list(wsdn.keys()):
                    if key in PRINT_NAMES_TO_CLEAN:
                        try:
                            del wsdn[key]
                            removed += 1
                        except Exception:
                            pass
            except Exception:
                continue

        logger.info(f"Удалено старых настроек печати: {removed}")
        return True
    except Exception as e:
        logger.warning(f"Не удалось очистить старые настройки: {e}")
        return False


def capture_original_setup(ws):
    """
    Снимает настройки страницы ДО их сброса.
    Вызывать строго до reset_page_setup_openpyxl — она их затирает.

    orientation может быть None: в файле ориентация не задана явно.
    Раньше такой случай молча превращался в «книжная», и файл намертво
    запирался в портрете.
    """
    def num(value, default):
        try:
            return float(value) if value is not None else default
        except Exception:
            return default

    ps = ws.page_setup
    m = ws.page_margins

    try:
        fit_to_page = bool(ws.sheet_properties.pageSetUpPr.fitToPage)
    except Exception:
        fit_to_page = False

    try:
        scale = int(ps.scale) if ps.scale else 100
    except Exception:
        scale = 100

    orientation_raw = str(getattr(ps, 'orientation', None) or '').lower()
    if orientation_raw == 'landscape':
        orientation = XL_LANDSCAPE
    elif orientation_raw == 'portrait':
        orientation = XL_PORTRAIT
    else:
        orientation = None          # в файле не задана — пусть решает форма

    return {
        'orientation': orientation,
        'scale': scale,
        'fit_to_page': fit_to_page,
        'left': num(m.left, 0.7), 'right': num(m.right, 0.7),
        'top': num(m.top, 0.75), 'bottom': num(m.bottom, 0.75),
        'header': num(m.header, 0.3), 'footer': num(m.footer, 0.3),
    }


def reset_page_setup_openpyxl(ws):
    """Сбрасывает настройки печати листа (без удаления содержимого)"""
    try:
        ws.print_area = None
        ws.print_title_rows = None
        ws.print_title_cols = None

        try:
            ws.page_setup.orientation = 'portrait'
        except Exception:
            pass

        # ГЛАВНЫЙ переключатель: пока fitToPage=True, Excel игнорирует scale
        try:
            if ws.sheet_properties.pageSetUpPr is None:
                ws.sheet_properties.pageSetUpPr = PageSetupProperties()
            ws.sheet_properties.pageSetUpPr.fitToPage = False
            ws.page_setup.fitToWidth = None
            ws.page_setup.fitToHeight = None
        except Exception as e:
            logger.debug(f"[{ws.title}] Не удалось снять fitToPage: {e}")

        try:
            ws.page_setup.scale = 100
        except Exception:
            pass

        ws.page_margins = PageMargins(
            left=MARGIN_LEFT, right=MARGIN_RIGHT,
            top=MARGIN_TOP, bottom=MARGIN_BOTTOM,
            header=MARGIN_HEADER, footer=MARGIN_FOOTER)

        try:
            if hasattr(ws, '_horizontal_page_breaks'):
                ws._horizontal_page_breaks = []
            if hasattr(ws, '_vertical_page_breaks'):
                ws._vertical_page_breaks = []
        except Exception:
            pass

        logger.info(f"[{ws.title}] Настройки печати сброшены (openpyxl)")
        return True
    except Exception as e:
        logger.warning(f"Ошибка сброса openpyxl: {e}")
        return False


def is_sheet_empty(ws):
    try:
        if ws.max_row <= 1 and ws.max_column <= 1:
            for row in ws.iter_rows(max_row=1, max_column=1):
                for cell in row:
                    if cell.value is not None and str(cell.value).strip():
                        return False
            return True
        return False
    except Exception:
        return True


def is_sheet_empty_com(ws):
    try:
        used = ws.UsedRange
        if used is None:
            return True
        val = getattr(used, "Value", None)
        if val is None:
            return True
        if not isinstance(val, tuple):
            return val in (None, "", 0)
        for row in val:
            if isinstance(row, tuple):
                for cell in row:
                    if cell not in (None, "", 0):
                        return False
            else:
                if row not in (None, "", 0):
                    return False
        return True
    except Exception:
        return True


def hard_clear_fills_only(wb):
    """Убирает только заливку во всех листах книги"""
    empty_fill = PatternFill(fill_type=None)
    for ws in wb.worksheets:
        try:
            if is_sheet_empty(ws):
                continue
            max_row = ws.max_row or 1
            max_col = ws.max_column or 1
            for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
                for cell in row:
                    try:
                        cell.fill = empty_fill
                    except Exception:
                        continue
            for rng in list(ws.merged_cells.ranges):
                try:
                    master = ws.cell(rng.min_row, rng.min_col)
                    master.fill = empty_fill
                except Exception:
                    continue
        except Exception as e:
            logger.warning(f"Ошибка очистки заливки на листе {ws.title}: {e}")
    logger.info("Удалена только заливка во всех листах книги")


# -----------------------
# Поиск шапки таблицы
# -----------------------
def find_header_rows(ws, first_data_row=1):
    """Поиск шапки по маркерам (№, п/п, А, Б, 1, 2, 3) и границам ячеек"""
    try:
        marker_row = None
        for col in [1, 2, 3]:
            for row in range(first_data_row, min(first_data_row + 25, ws.max_row + 1)):
                cell_value = ws.cell(row=row, column=col).value
                cell_str = str(cell_value).strip().lower() if cell_value else ""
                if cell_str in ['a', 'b', 'c', 'а', 'б', 'в', 'п/п', '№']:
                    marker_row = row
                    break
                elif re.fullmatch(r'^[123]$', cell_str):
                    marker_row = row
                    break
            if marker_row is not None:
                break

        if marker_row is None:
            logger.info(f"[{ws.title}] Маркер не найден → переходим к поиску по границам.")
            return find_header_rows_any_border_advanced(ws)

        header_block = [marker_row]
        current_row = marker_row - 1
        while current_row >= 1:
            has_right_border = False
            for col in range(1, min(12, ws.max_column + 1)):
                border = ws.cell(row=current_row, column=col).border
                if border.right and border.right.style:
                    has_right_border = True
                    break
            if has_right_border:
                header_block.append(current_row)
                current_row -= 1
            else:
                break

        start_row, end_row = min(header_block), max(header_block)
        logger.info(f"[{ws.title}] ✅ Заголовок найден: строки {start_row}-{end_row}")
        return start_row, end_row
    except Exception as e:
        logger.error(f"[{ws.title}] Ошибка при поиске заголовков: {e}")
        return 1, 1


def find_header_rows_any_border_advanced(ws, max_search_rows=30):
    """Резервный поиск шапки по первому непрерывному блоку строк с границами"""
    try:
        all_border_rows = []
        for row in range(1, min(max_search_rows, ws.max_row + 1)):
            for col in range(1, min(20, ws.max_column + 1)):
                border = ws.cell(row=row, column=col).border
                if (border.top and border.top.style) or (border.right and border.right.style):
                    all_border_rows.append(row)
                    break

        if not all_border_rows:
            logger.info(f"[{ws.title}] ⚠️ Границы не найдены — используем первую строку.")
            return 1, 1

        blocks = []
        current_block = [all_border_rows[0]]
        for i in range(1, len(all_border_rows)):
            if all_border_rows[i] == all_border_rows[i - 1] + 1:
                current_block.append(all_border_rows[i])
            else:
                blocks.append(current_block)
                current_block = [all_border_rows[i]]
        blocks.append(current_block)

        first_block = blocks[0]
        start_row, end_row = min(first_block), max(first_block)
        logger.info(f"[{ws.title}] 🧩 Границы найдены: {start_row}-{end_row}")
        return start_row, end_row
    except Exception as e:
        logger.error(f"[{ws.title}] Ошибка поиска по границам: {e}")
        return 1, 1


def find_header_rows_smart(ws):
    try:
        start_row, end_row = find_header_rows(ws)
        if end_row - start_row > 8:
            logger.warning(f"[{ws.title}] Заголовок слишком большой ({start_row}:{end_row}) → 2–4")
            start_row, end_row = 2, min(4, ws.max_row)
        return start_row, end_row
    except Exception as e:
        logger.error(f"[{ws.title}] Ошибка при комбинированном поиске: {e}")
        return 1, 1


def safe_add_page_number(ws):
    """Добавляет номер страницы в нижний колонтитул"""
    try:
        try:
            ws.oddFooter.center.text = "Страница &P из &N"
        except Exception:
            try:
                ws.oddFooter.center = "Страница &P из &N"
            except Exception:
                pass
        try:
            ws.evenFooter.center.text = "Страница &P из &N"
        except Exception:
            try:
                ws.evenFooter.center = "Страница &P из &N"
            except Exception:
                pass
        return True
    except Exception as e:
        logger.warning(f"Ошибка добавления номера страницы: {e}")
        return False


# -----------------------
# Расчёт области печати (openpyxl — только расчёт!)
# -----------------------
def is_real_data_cell(ws, row, col, merged_ranges=None):
    """Строгая проверка: является ли ячейка реальными данными"""
    try:
        cell = ws.cell(row=row, column=col)
        if cell.value is not None:
            cell_str = str(cell.value).strip()
            if cell_str and cell_str not in ['', ' ', '\t', '\n', '\r']:
                return True

        ranges = merged_ranges if merged_ranges is not None else ws.merged_cells.ranges
        for merged_range in ranges:
            if (merged_range.min_row <= row <= merged_range.max_row and
                    merged_range.min_col <= col <= merged_range.max_col):
                master_cell = ws.cell(merged_range.min_row, merged_range.min_col)
                if master_cell.value is not None:
                    return True
        return False
    except Exception:
        return False


def calc_print_area(ws):
    """
    Считает границы области печати по реальным данным, возвращает "A1:K1036".

    Потолок сканирования MAX_SCAN_ROWS большой намеренно: жёсткий лимит
    обрезал бы живые данные (реальный случай — лист на 1034 строки при
    лимите 1000). От бесконечных листов защищает ранняя остановка
    после 5 пустых строк подряд.
    """
    try:
        merged_ranges = list(ws.merged_cells.ranges)
        max_real_row = 1
        max_real_col = 1
        empty_block_count = 0

        for row in range(1, min(ws.max_row, MAX_SCAN_ROWS) + 1):
            row_has_real_data = False
            for col in range(1, min(ws.max_column, MAX_SCAN_COLS) + 1):
                if is_real_data_cell(ws, row, col, merged_ranges):
                    row_has_real_data = True
                    max_real_row = max(max_real_row, row)
                    max_real_col = max(max_real_col, col)

            if not row_has_real_data and row > 10:
                empty_block_count += 1
            else:
                empty_block_count = 0
            if empty_block_count >= 5:
                break

        print_max_row = min(max_real_row + PRINT_AREA_ROW_PADDING, ws.max_row)
        print_max_col = max_real_col          # по колонкам запаса НЕТ

        area = f"A1:{get_column_letter(print_max_col)}{print_max_row}"
        try:
            ws.print_area = area
        except Exception:
            pass
        logger.info(f"[{ws.title}] Область печати рассчитана: {area} "
                    f"(данные до {get_column_letter(max_real_col)}{max_real_row}, "
                    f"лист тянется до {get_column_letter(ws.max_column)}{ws.max_row})")
        return area
    except Exception as e:
        logger.warning(f"Ошибка расчёта области печати: {e}")
        return None


def aggressive_clean_empty_regions(ws):
    """Снимает пустые объединения и убирает случайные границы за пределами данных"""
    try:
        last_real_row = 1
        last_real_col = 1

        # ВАЖНО: list() — иначе "Set changed size during iteration" при unmerge
        for merged_range in list(ws.merged_cells.ranges):
            master_cell = ws.cell(merged_range.min_row, merged_range.min_col)
            if master_cell.value is None:
                try:
                    ws.unmerge_cells(str(merged_range))
                except Exception:
                    pass
            else:
                last_real_row = max(last_real_row, merged_range.max_row)
                last_real_col = max(last_real_col, merged_range.max_col)

        merged_ranges = list(ws.merged_cells.ranges)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                try:
                    if is_real_data_cell(ws, cell.row, cell.column, merged_ranges):
                        last_real_row = max(last_real_row, cell.row)
                        last_real_col = max(last_real_col, cell.column)
                except Exception:
                    continue

        empty_border = Border()
        for row in range(last_real_row + 5, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                try:
                    cell = ws.cell(row=row, column=col)
                    border = cell.border
                    has_border = ((border.top and border.top.style) or
                                  (border.bottom and border.bottom.style) or
                                  (border.left and border.left.style) or
                                  (border.right and border.right.style))
                    if has_border and cell.value is None:
                        cell.border = empty_border
                except Exception:
                    continue

        return True
    except Exception as e:
        logger.warning(f"Ошибка агрессивной очистки: {e}")
        return False


# -----------------------
# Диагностика
# -----------------------
def log_com_pagesetup(ws):
    try:
        ps = ws.PageSetup
        logger.info(f"[{ws.Name}] Итог: Zoom={getattr(ps, 'Zoom', None)}, "
                    f"Orientation={getattr(ps, 'Orientation', None)}, "
                    f"PrintArea={getattr(ps, 'PrintArea', None)}")
    except Exception:
        pass


# -----------------------
# COM / Excel
# -----------------------
def create_excel_instance(visible=False):
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = visible
    excel.DisplayAlerts = False
    return excel


def find_last_cell_com(ws):
    """Резерв: спрашиваем у Excel последнюю ячейку с данными"""
    try:
        first_cell = ws.Cells(1, 1)
        last_row_cell = ws.Cells.Find("*", first_cell, XL_FORMULAS, None,
                                      XL_BY_ROWS, XL_PREVIOUS)
        last_col_cell = ws.Cells.Find("*", first_cell, XL_FORMULAS, None,
                                      XL_BY_COLUMNS, XL_PREVIOUS)
        if last_row_cell is None or last_col_cell is None:
            return None, None
        return int(last_row_cell.Row), int(last_col_cell.Column)
    except Exception as e:
        logger.debug(f"[{ws.Name}] Cells.Find недоступен: {e}")
        return None, None


def set_print_area_com(ws, preferred=None):
    """
    Ставит область печати средствами COM и ПРОВЕРЯЕТ результат чтением обратно.
      1) значение, рассчитанное openpyxl (например "A1:K1036")
      2) резерв — Cells.Find у самого Excel
    """
    ps = ws.PageSetup

    if preferred:
        try:
            ps.PrintArea = preferred
            applied = ps.PrintArea
            if applied:
                logger.info(f"[{ws.Name}] Область печати установлена: {applied}")
                return applied
            logger.warning(f"[{ws.Name}] PrintArea='{preferred}' не применилась, пробуем Find")
        except Exception as e:
            logger.warning(f"[{ws.Name}] Не удалось задать PrintArea='{preferred}': {e}")

    last_row, last_col = find_last_cell_com(ws)
    if last_row and last_col:
        try:
            area = ws.Range(ws.Cells(1, 1), ws.Cells(last_row, last_col)).Address
            ps.PrintArea = area
            applied = ps.PrintArea
            logger.info(f"[{ws.Name}] Область печати через Find: {applied}")
            return applied
        except Exception as e:
            logger.warning(f"[{ws.Name}] Find-вариант не сработал: {e}")

    logger.warning(f"[{ws.Name}] Область печати НЕ задана — ширина будет "
                   f"меряться по UsedRange, возможен заниженный масштаб")
    return None


def get_print_range_com(ws):
    """Range области печати (или UsedRange, если PrintArea не задана)"""
    try:
        pa = ws.PageSetup.PrintArea
        if pa:
            return ws.Range(pa.split('!')[-1])
    except Exception:
        pass
    return ws.UsedRange


def apply_our_margins_com(ws, excel):
    """Ставит наши поля страницы"""
    ps = ws.PageSetup
    ps.LeftMargin = excel.InchesToPoints(MARGIN_LEFT)
    ps.RightMargin = excel.InchesToPoints(MARGIN_RIGHT)
    ps.TopMargin = excel.InchesToPoints(MARGIN_TOP)
    ps.BottomMargin = excel.InchesToPoints(MARGIN_BOTTOM)
    ps.HeaderMargin = excel.InchesToPoints(MARGIN_HEADER)
    ps.FooterMargin = excel.InchesToPoints(MARGIN_FOOTER)


def preserve_if_fits_com(ws, excel, orig):
    """
    Если с РОДНЫМИ полями, ориентацией и масштабом таблица помещается
    по ширине — восстанавливает их и возвращает True.
    Иначе ничего не меняет и возвращает False.

    PRINTER_MARGIN_PTS здесь НЕ вычитается: файл уже существует и печатался
    в таком виде, поправка на принтер в нём учтена автором. Погрешность
    измерения ширины покрывает PRESERVE_TOLERANCE.
    """
    if not PRESERVE_IF_FITS or not orig:
        return False

    # ориентация в файле не задана — сохранять нечего, пусть решает форма
    orientation = orig.get('orientation')
    if orientation is None:
        logger.info(f"[{ws.Name}] Ориентация в оригинале не задана — переформатируем")
        return False

    # «Вписать в страницу» — масштаб в файле фиктивный, доверять нечему
    if orig.get('fit_to_page'):
        logger.info(f"[{ws.Name}] В оригинале включено «вписать в страницу» — "
                    f"переформатируем")
        return False

    scale = orig.get('scale', 100)
    if scale < MIN_SCALE or scale > 100:
        logger.info(f"[{ws.Name}] Исходный масштаб {scale}% вне диапазона "
                    f"{MIN_SCALE}–100% — переформатируем")
        return False

    try:
        data_w = float(get_print_range_com(ws).Width or 0.0)
    except Exception:
        return False
    if data_w <= 0:
        return False

    page_w_inches = (PAGE_WIDTH_A4_PORTRAIT if orientation == XL_PORTRAIT
                     else PAGE_HEIGHT_A4_PORTRAIT)
    available = excel.InchesToPoints(
        page_w_inches - orig.get('left', 0.7) - orig.get('right', 0.7))
    needed = data_w * scale / 100.0

    if needed > available * PRESERVE_TOLERANCE:
        logger.info(f"[{ws.Name}] Оригинал НЕ помещается "
                    f"(нужно {needed:.0f} pts, доступно {available:.0f} pts) — "
                    f"переформатируем")
        return False

    # помещается — восстанавливаем всё как было
    try:
        ps = ws.PageSetup
        ps.PaperSize = XL_PAPER_A4
        ps.Orientation = orientation
        ps.LeftMargin = excel.InchesToPoints(orig.get('left', 0.7))
        ps.RightMargin = excel.InchesToPoints(orig.get('right', 0.7))
        ps.TopMargin = excel.InchesToPoints(orig.get('top', 0.75))
        ps.BottomMargin = excel.InchesToPoints(orig.get('bottom', 0.75))
        ps.HeaderMargin = excel.InchesToPoints(orig.get('header', 0.3))
        ps.FooterMargin = excel.InchesToPoints(orig.get('footer', 0.3))
        ps.FitToPagesWide = False
        ps.FitToPagesTall = False
        ps.Zoom = scale

        logger.info(f"[{ws.Name}] ОРИГИНАЛ ПОМЕЩАЕТСЯ — оставляем как есть: "
                    f"{'книжная' if orientation == XL_PORTRAIT else 'альбомная'}, "
                    f"zoom={scale}%, поля {orig.get('left'):.2f}/{orig.get('right'):.2f}\" "
                    f"(нужно {needed:.0f} pts, доступно {available:.0f} pts)")
        return True
    except Exception as e:
        logger.warning(f"[{ws.Name}] Не удалось восстановить исходные настройки: {e}")
        return False


def wrap_text_com(ws):
    """Включает перенос по словам в области печати"""
    if not ENABLE_WRAP_TEXT:
        return False
    try:
        get_print_range_com(ws).WrapText = True
        logger.info(f"[{ws.Name}] Перенос по словам включён")
        return True
    except Exception as e:
        logger.warning(f"[{ws.Name}] Не удалось включить перенос: {e}")
        return False


def measure_wrapped_height(helper_ws, text, width_pts, font):
    """
    Измеряет, сколько нужно высоты (pts) для текста при заданной ширине.

    Считаем в отдельной служебной книге: подгоняем ширину колонки под нужную,
    включаем перенос, вызываем AutoFit и читаем получившуюся высоту.
    Это единственный надёжный способ — для объединённых ячеек Excel
    AutoFit не выполняет вообще.
    """
    try:
        col = helper_ws.Columns(1)
        col.ColumnWidth = 20.0
        base_pts = float(col.Width or 0.0)
        if base_pts <= 0:
            return 0.0
        target_cw = 20.0 * width_pts / base_pts
        col.ColumnWidth = max(1.0, min(255.0, target_cw))

        cell = helper_ws.Cells(1, 1)
        cell.ClearContents()
        cell.NumberFormat = "@"          # чтобы текст не стал формулой/датой
        try:
            cell.Font.Name = font.Name
            cell.Font.Size = font.Size
            cell.Font.Bold = font.Bold
            cell.Font.Italic = font.Italic
        except Exception:
            pass
        cell.WrapText = True
        cell.Value = text
        helper_ws.Rows(1).AutoFit()
        return float(helper_ws.Rows(1).RowHeight or 0.0)
    except Exception as e:
        logger.debug(f"Не удалось измерить высоту текста: {e}")
        return 0.0


def fix_header_heights_com(ws, start_row, end_row, last_col, helper_ws):
    """
    Увеличивает высоту строк шапки там, где перенесённый текст не помещается.

    Rows.AutoFit() тут бесполезен: объединённые ячейки он игнорирует.
    Высоту ТОЛЬКО увеличиваем — вручную подогнанные значения не портим.
    """
    if not FIX_HEADER_HEIGHTS:
        return 0
    grown = 0
    try:
        for r in range(start_row, end_row + 1):
            for c in range(1, last_col + 1):
                try:
                    cell = ws.Cells(r, c)
                    val = cell.Value
                    if val is None or not str(val).strip():
                        continue
                    text = str(val)

                    if cell.MergeCells:
                        area = cell.MergeArea
                        # работаем только с главной ячейкой блока
                        if int(area.Row) != r or int(area.Column) != c:
                            continue
                        width_pts = float(area.Width or 0.0)
                        height_pts = float(area.Height or 0.0)
                        target_row = int(area.Row) + int(area.Rows.Count) - 1
                    else:
                        width_pts = float(cell.Width or 0.0)
                        height_pts = float(ws.Rows(r).RowHeight or 0.0)
                        target_row = r

                    if width_pts <= 0:
                        continue

                    needed = measure_wrapped_height(helper_ws, text, width_pts, cell.Font)
                    if needed <= 0:
                        continue

                    if needed > height_pts + HEIGHT_TOLERANCE_PTS:
                        delta = needed - height_pts
                        current = float(ws.Rows(target_row).RowHeight or 0.0)
                        new_height = min(current + delta, MAX_HEADER_ROW_HEIGHT)
                        if new_height > current + 0.5:
                            ws.Rows(target_row).RowHeight = new_height
                            grown += 1
                            logger.info(
                                f"[{ws.Name}] {cell.Address(False, False)}: высота "
                                f"{height_pts:.0f}→{height_pts + (new_height - current):.0f} pt "
                                f"(нужно {needed:.0f}), строка {target_row}")
                except Exception:
                    continue

        if grown:
            logger.info(f"[{ws.Name}] Подогнано высот в шапке: {grown}")
        return grown
    except Exception as e:
        logger.warning(f"[{ws.Name}] Ошибка подгонки высоты шапки: {e}")
        return grown


# -----------------------
# Выбор ориентации по форме таблицы
# -----------------------
def decide_layout(p, l, data_w, data_h, min_scale):
    """
    Возвращает (ориентация, масштаб или None, причина).
    Масштаб None означает "не подбирать под страницу".

    p, l      — масштабы, при которых таблица влезает по ширине
                в книжной и альбомной соответственно
    data_w/h  — ширина и высота области печати в pts
    """
    tall = TALL_TABLE_ORIENTATION if TALL_TABLE_ORIENTATION else XL_PORTRAIT

    # Не помещается по ширине даже в альбомной при минимальном масштабе.
    # Это широкая таблица. Переворачивать в книжную бессмысленно:
    # книжная страница уже, масштаб получится ещё мельче.
    if l < min_scale:
        return XL_LANDSCAPE, None, (
            "широкая таблица, на страницу не помещается "
            "(портрет %d%%, альбом %d%%) — альбомная, печать лентой" % (p, l))

    # По ширине проходит и в книжной, и в альбомной — решает форма.
    if p >= min_scale:
        if data_h >= data_w * TALL_TABLE_RATIO:
            zoom = p if tall == XL_PORTRAIT else l
            return tall, zoom, (
                "вытянута в высоту (%.0fx%.0f pts)" % (data_h, data_w))
        return XL_LANDSCAPE, l, (
            "вытянута в ширину (%.0fx%.0f pts)" % (data_w, data_h))

    # В книжную по ширине не влезает, в альбомную влезает.
    return XL_LANDSCAPE, l, "по ширине проходит только в альбомной"


def fit_sheet_com(ws, excel, orig=None, min_scale=None):
    """
    Подбирает ориентацию и масштаб по РЕАЛЬНЫМ размерам области печати.
    Ориентация выбирается по форме таблицы (см. decide_layout).

    orig — словарь из capture_original_setup. Нужен только чтобы вернуть
    авторский масштаб широкой таблице, которая на страницу не влезает.
    """
    if min_scale is None:
        min_scale = MIN_SCALE

    ps = ws.PageSetup
    ps.PaperSize = XL_PAPER_A4
    ps.FitToPagesWide = False
    ps.FitToPagesTall = False
    ps.Zoom = 100

    rng = get_print_range_com(ws)
    data_w = float(rng.Width or 0.0)
    data_h = float(rng.Height or 0.0)
    if data_w <= 0:
        logger.warning(f"[{ws.Name}] Нулевая ширина — масштаб не меняем")
        return None

    def scale_for(orientation):
        ps.Orientation = orientation
        page_w = excel.InchesToPoints(
            PAGE_WIDTH_A4_PORTRAIT if orientation == XL_PORTRAIT
            else PAGE_HEIGHT_A4_PORTRAIT)
        avail = page_w - ps.LeftMargin - ps.RightMargin - PRINTER_MARGIN_PTS
        if avail <= 0:
            return 1
        return max(1, min(100, int(avail / data_w * 100)))

    p = scale_for(XL_PORTRAIT)
    l = scale_for(XL_LANDSCAPE)

    orient, zoom, reason = decide_layout(p, l, data_w, data_h, min_scale)
    ps.Orientation = orient
    name = "книжная" if orient == XL_PORTRAIT else "альбомная"

    # --- широкая таблица: масштаб под страницу не подбираем ---
    if zoom is None:
        applied = 100
        if WIDE_TABLE_KEEP_ORIGINAL_SCALE and orig:
            src_scale = orig.get('scale', 100)
            if not orig.get('fit_to_page') and min_scale <= src_scale <= 100:
                applied = src_scale
        try:
            ps.Zoom = applied
        except Exception as e:
            logger.debug(f"[{ws.Name}] Не удалось задать Zoom={applied}: {e}")
        logger.info(f"[{ws.Name}] {name} ({reason}), zoom={ps.Zoom}% — "
                    f"под страницу не ужимаем")
        return ps.Zoom

    # --- обычный случай: подбираем масштаб ---
    ps.Zoom = zoom
    calculated_zoom = zoom

    # Контрольная проверка у Excel. Если VPageBreaks не обнуляется —
    # проверке НЕ доверяем и возвращаем расчётное значение,
    # иначе масштаб молча уезжает в пол MIN_SCALE.
    try:
        ws.Parent.Activate()
        ws.Activate()
        guard = 0
        confirmed = False
        while guard < 15:
            if ws.VPageBreaks.Count == 0:
                confirmed = True
                break
            if ps.Zoom <= min_scale:
                break
            ps.Zoom = ps.Zoom - 1
            guard += 1
        if not confirmed:
            ps.Zoom = calculated_zoom
            logger.debug(f"[{ws.Name}] VPageBreaks не обнулился — "
                         f"оставляем {calculated_zoom}%")
    except Exception as e:
        logger.debug(f"[{ws.Name}] VPageBreaks недоступны: {e}")
        ps.Zoom = calculated_zoom

    logger.info(f"[{ws.Name}] {name} ({reason}), zoom={ps.Zoom}% "
                f"(расчёт: портрет {p}%, альбом {l}%, ширина {data_w:.0f} pts)")
    return ps.Zoom


# -----------------------
# Основная обработка файла
# -----------------------
def get_modified_path(original_path, root_folder):
    original_path = Path(original_path)
    root_folder = Path(root_folder)
    try:
        if root_folder not in original_path.parents:
            rel_path = Path(original_path.name)
        else:
            rel_path = original_path.relative_to(root_folder)
    except ValueError:
        rel_path = Path(original_path.name)
    return original_path.parent / "modified" / rel_path


def process_excel_file(file_path, log_func=None, excel_app=None):
    """openpyxl → save → COM финализирует PageSetup. Возвращает True/False."""
    temp_xlsx = None
    wb = None
    local_excel = False
    header_rows_map = {}
    print_areas = {}
    orig_setups = {}
    excel = None

    try:
        file_path = str(file_path)
        original_path_for_modified = Path(file_path)
        if not os.path.exists(file_path):
            logger.info(f"Файл не существует: {file_path}")
            return False

        if excel_app is None:
            excel = create_excel_instance(visible=False)
            local_excel = True
        else:
            excel = excel_app

        # --- Конвертация .xls -> .xlsx ---
        if file_path.lower().endswith('.xls') and not file_path.lower().endswith('.xlsx'):
            temp_xlsx = file_path + 'x'
            try:
                logger.info(f"Конвертируем {file_path} -> {temp_xlsx}")
                wb_conv = None
                try:
                    wb_conv = excel.Workbooks.Open(file_path)
                    wb_conv.SaveAs(temp_xlsx, FileFormat=51)
                finally:
                    if wb_conv is not None:
                        try:
                            wb_conv.Close()
                        except Exception:
                            pass
                file_path_for_openpyxl = temp_xlsx
            except Exception as e:
                logger.info(f"Ошибка конвертации: {e}")
                return False
        else:
            file_path_for_openpyxl = file_path

        # === OPENPYXL ОБРАБОТКА ===
        try:
            wb = load_workbook(filename=file_path_for_openpyxl, read_only=False, keep_vba=False)
            cleanup_old_print_settings(wb)
            hard_clear_fills_only(wb)

            modified_path = get_modified_path(
                original_path_for_modified,
                os.path.dirname(original_path_for_modified)).with_suffix('.xlsx')
            modified_path.parent.mkdir(parents=True, exist_ok=True)

            for ws in wb.worksheets:
                try:
                    if is_sheet_empty(ws):
                        logger.info(f"[{ws.title}] Пустой лист — пропускаем")
                        continue

                    # СНИМАЕМ настройки оригинала ДО сброса — reset их затирает
                    orig = capture_original_setup(ws)
                    orig_setups[ws.title] = orig
                    if orig['orientation'] is None:
                        orient_name = 'не задана'
                    elif orig['orientation'] == XL_PORTRAIT:
                        orient_name = 'книжная'
                    else:
                        orient_name = 'альбомная'
                    logger.info(
                        f"[{ws.title}] Оригинал: {orient_name}, "
                        f"scale={orig['scale']}%, fitToPage={orig['fit_to_page']}, "
                        f"поля {orig['left']:.2f}/{orig['right']:.2f}\"")

                    reset_page_setup_openpyxl(ws)

                    start_row, end_row = find_header_rows_smart(ws)
                    if end_row - start_row > 5:
                        logger.warning(f"[{ws.title}] Слишком большой заголовок "
                                       f"{start_row}:{end_row}, используем 1:2")
                        start_row, end_row = 1, 2
                    header_rows_map[ws.title] = (start_row, end_row)
                    ws.print_title_rows = f"{start_row}:{end_row}"
                    ws.print_title_cols = "A:A"

                    if safe_add_page_number(ws):
                        logger.info(f"[{ws.title}] Добавлен номер страницы")

                    try:
                        ws.page_setup.paperSize = ws.PAPERSIZE_A4
                    except Exception:
                        pass

                    aggressive_clean_empty_regions(ws)

                    area = calc_print_area(ws)
                    if area:
                        print_areas[ws.title] = area

                except Exception as e:
                    logger.info(f"Ошибка обработки листа {ws.title}: {e}")
                    continue

            wb.save(str(modified_path))
            logger.info(f"Сохранено после openpyxl обработки: {modified_path}")

        except Exception as e:
            logger.info(f"Ошибка загрузки/обработки openpyxl: {e}")
            return False

        # === COM ФИНАЛИЗАЦИЯ ===
        wb_com = None
        helper_wb = None
        helper_ws = None
        try:
            wb_com = excel.Workbooks.Open(str(modified_path))

            # служебная книга для измерения высоты текста
            if FIX_HEADER_HEIGHTS:
                try:
                    helper_wb = excel.Workbooks.Add()
                    helper_ws = helper_wb.Worksheets(1)
                except Exception as e:
                    logger.warning(f"Не удалось создать служебную книгу: {e}")
                    helper_ws = None

            for ws in wb_com.Sheets:
                try:
                    if is_sheet_empty_com(ws):
                        logger.info(f"[{ws.Name}] Пустой лист (COM) — пропускаем")
                        continue

                    ps = ws.PageSetup
                    orig = orig_setups.get(ws.Name)

                    # 1) область печати — ДО любых замеров ширины
                    set_print_area_com(ws, print_areas.get(ws.Name))

                    # 2) сквозные строки и столбец (ставим всегда — это суть задачи)
                    if ws.Name in header_rows_map:
                        start_row, end_row = header_rows_map[ws.Name]
                        try:
                            ps.PrintTitleRows = f"${start_row}:${end_row}"
                            ps.PrintTitleColumns = "$A:$A"
                            logger.info(f"[{ws.Name}] Сквозные строки: {start_row}:{end_row}")
                        except Exception:
                            pass

                    # 3) оригинал уже помещается? тогда не трогаем геометрию
                    preserved = preserve_if_fits_com(ws, excel, orig)

                    if not preserved:
                        apply_our_margins_com(ws, excel)

                    # 4) перенос текста и высота шапки
                    if not preserved or WRAP_TEXT_WHEN_PRESERVED:
                        wrap_text_com(ws)
                        if helper_ws is not None and ws.Name in header_rows_map:
                            start_row, end_row = header_rows_map[ws.Name]
                            try:
                                last_col = int(get_print_range_com(ws).Columns.Count)
                            except Exception:
                                last_col = 1
                            fix_header_heights_com(ws, start_row, end_row,
                                                   last_col, helper_ws)

                    # 5) ориентация и масштаб — только если переформатируем
                    if not preserved:
                        fit_sheet_com(ws, excel, orig)

                    log_com_pagesetup(ws)

                    try:
                        ps.Order = 2      # xlOverThenDown
                    except Exception:
                        pass

                except Exception as e:
                    logger.info(f"[{ws.Name}] Ошибка COM настроек: {e}")

            wb_com.Save()
            logger.info(f"COM-финализация завершена: {modified_path}")

        except Exception as e:
            logger.info(f"Win32COM ошибка: {e}")
            return False
        finally:
            if helper_wb is not None:
                try:
                    helper_wb.Close(SaveChanges=False)
                except Exception:
                    pass
            if wb_com is not None:
                try:
                    wb_com.Close()
                except Exception:
                    pass

        return True

    finally:
        if wb:
            try:
                wb.close()
            except Exception:
                pass
        if temp_xlsx and os.path.exists(temp_xlsx):
            try:
                os.remove(temp_xlsx)
                logger.info(f"Удалён временный файл: {temp_xlsx}")
            except Exception as e:
                logger.info(f"Не удалось удалить временный файл: {e}")
        if local_excel and excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass


# -----------------------
# GUI
# -----------------------
class ExcelProcessor:
    def __init__(self):
        self.root = tk.Tk()
        self.root.withdraw()
        self.folders = []                      # список выбранных папок (Path)
        self.progress_var = tk.DoubleVar()
        self.status_var = tk.StringVar(value="Готово")
        self.delete_archives_var = tk.BooleanVar(value=False)
        self.progress_bar = None
        self.log_text = None
        self.folder_list = None
        self.last_dir = None                   # чтобы диалог открывался там же
        self.setup_ui()

    def setup_ui(self):
        self.root.deiconify()
        self.root.title("Wizard Archive")
        self.root.geometry("760x620")

        main_frame = ttk.Frame(self.root, padding="15")
        main_frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(main_frame, text="Wizard Archive",
                  font=('Helvetica', 14, 'bold')).pack(pady=(0, 10))

        # --- список папок ---
        folders_frame = ttk.LabelFrame(main_frame, text="Папки для обработки", padding="10")
        folders_frame.pack(fill=tk.BOTH, expand=True)

        list_frame = ttk.Frame(folders_frame)
        list_frame.pack(fill=tk.BOTH, expand=True)

        # EXTENDED = Ctrl+клик и Shift+клик для выделения нескольких строк
        self.folder_list = tk.Listbox(list_frame, height=6, selectmode=tk.EXTENDED)
        self.folder_list.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        list_scroll = ttk.Scrollbar(list_frame, command=self.folder_list.yview)
        list_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.folder_list.config(yscrollcommand=list_scroll.set)

        # Delete удаляет выделенные папки
        self.folder_list.bind("<Delete>", lambda e: self.remove_selected_folders())

        folder_btns = ttk.Frame(folders_frame)
        folder_btns.pack(fill="x", pady=(8, 0))
        ttk.Button(folder_btns, text="Добавить папку",
                   command=self.add_folder).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(folder_btns, text="Удалить выбранные",
                   command=self.remove_selected_folders).pack(side=tk.LEFT, padx=5)
        ttk.Button(folder_btns, text="Очистить список",
                   command=self.clear_folders).pack(side=tk.LEFT, padx=5)
        ttk.Label(folder_btns,
                  text="Ctrl+клик / Shift+клик — выделить несколько, Delete — удалить",
                  foreground="gray").pack(side=tk.LEFT, padx=10)

        # --- опции и запуск ---
        ttk.Checkbutton(main_frame, text="Удалить архивы после распаковки",
                        variable=self.delete_archives_var).pack(anchor="w", pady=(10, 5))

        button_frame = ttk.Frame(main_frame)
        button_frame.pack(pady=5)
        ttk.Button(button_frame, text="Только распаковать архивы",
                   command=self.unpack_only).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="Полная обработка",
                   command=self.start_processing).pack(side=tk.LEFT, padx=5)

        self.progress_bar = ttk.Progressbar(main_frame, variable=self.progress_var, maximum=100)
        self.progress_bar.pack(fill="x", pady=10)

        status_frame = ttk.Frame(main_frame)
        status_frame.pack(fill="x", pady=5)
        ttk.Label(status_frame, text="Статус:").pack(side=tk.LEFT)
        ttk.Label(status_frame, textvariable=self.status_var).pack(side=tk.LEFT, padx=10)

        log_frame = ttk.LabelFrame(main_frame, text="Лог выполнения", padding="10")
        log_frame.pack(fill=tk.BOTH, expand=True, pady=(5, 0))
        self.log_text = tk.Text(log_frame, height=10, width=80, wrap=tk.WORD)
        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar = ttk.Scrollbar(log_frame, command=self.log_text.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.log_text.config(yscrollcommand=scrollbar.set)

        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    # -----------------------
    # Работа со списком папок
    # -----------------------
    def add_folder(self):
        """Системный диалог не умеет выбирать несколько папок сразу,
        поэтому набираем список по одной."""
        folder = filedialog.askdirectory(title="Выберите папку",
                                         initialdir=self.last_dir or os.path.expanduser("~"))
        if not folder:
            return
        path = Path(folder)
        self.last_dir = str(path.parent)

        if path in self.folders:
            self.log_to_gui(f"Папка уже в списке: {path.name}")
            return

        # предупреждаем о вложенности — иначе файлы обработаются дважды
        for existing in self.folders:
            if existing in path.parents:
                if not messagebox.askyesno(
                        "Вложенная папка",
                        f"'{path.name}' находится внутри уже выбранной "
                        f"'{existing.name}'.\nЕё файлы будут обработаны дважды.\n\n"
                        f"Всё равно добавить?"):
                    return
            elif path in existing.parents:
                if not messagebox.askyesno(
                        "Вложенная папка",
                        f"Уже выбранная '{existing.name}' находится внутри "
                        f"'{path.name}'.\nЕё файлы будут обработаны дважды.\n\n"
                        f"Всё равно добавить?"):
                    return

        self.folders.append(path)
        self.folder_list.insert(tk.END, str(path))
        self.log_to_gui(f"Добавлена папка: {path}")

    def remove_selected_folders(self):
        selected = list(self.folder_list.curselection())
        if not selected:
            messagebox.showinfo("Ничего не выбрано",
                                "Выделите папки в списке (Ctrl+клик для нескольких)")
            return
        for index in reversed(selected):       # с конца, чтобы индексы не съезжали
            removed = self.folders.pop(index)
            self.folder_list.delete(index)
            self.log_to_gui(f"Убрана папка: {removed.name}")

    def clear_folders(self):
        self.folders.clear()
        self.folder_list.delete(0, tk.END)
        self.log_to_gui("Список папок очищен")

    def get_folders(self):
        """Проверяет, что папки выбраны и существуют"""
        if not self.folders:
            messagebox.showerror("Ошибка", "Сначала добавьте хотя бы одну папку!")
            return None
        missing = [f for f in self.folders if not f.exists()]
        if missing:
            messagebox.showerror("Ошибка", "Папки не найдены:\n" +
                                 "\n".join(str(m) for m in missing))
            return None
        return list(self.folders)

    # -----------------------
    # Логирование
    # -----------------------
    def log_to_gui(self, message):
        """Потокобезопасная запись в лог: работа идёт из фонового потока"""
        def _append():
            timestamp = time.strftime("%H:%M:%S")
            self.log_text.insert(tk.END, f"[{timestamp}] {message}\n")
            self.log_text.see(tk.END)
        try:
            self.root.after(0, _append)
        except Exception:
            pass
        print(message)

    def set_status(self, text):
        try:
            self.root.after(0, lambda: self.status_var.set(text))
        except Exception:
            pass

    def set_progress(self, value):
        try:
            self.root.after(0, lambda: self.progress_var.set(value))
        except Exception:
            pass

    # -----------------------
    # Распаковка
    # -----------------------
    def unpack_only(self):
        folders = self.get_folders()
        if not folders:
            return
        import threading
        threading.Thread(target=self._unpack_thread, args=(folders,), daemon=True).start()

    def _unpack_thread(self, folders):
        self.set_status("Распаковка архивов...")
        self.log_to_gui(f"Начинаю распаковку архивов ({len(folders)} папок)...")
        total_archives = 0
        try:
            for i, folder_path in enumerate(folders, 1):
                self.set_status(f"Распаковка: {folder_path.name} ({i}/{len(folders)})")
                self.set_progress(i * 100 / len(folders))
                self.log_to_gui(f"--- Папка: {folder_path}")

                processed = process_archives_in_folder(folder_path,
                                                       self.delete_archives_var.get())
                total_archives += len(processed)
                deleted_count = clean_empty_folders(folder_path)
                if deleted_count > 0:
                    self.log_to_gui(f"Удалено {deleted_count} пустых папок")

            self.set_status(f"Распаковано архивов: {total_archives}")
            self.set_progress(100)
            if total_archives:
                delete_status = "УДАЛЕНЫ" if self.delete_archives_var.get() else "СОХРАНЕНЫ"
                messagebox.showinfo("Готово",
                                    f"Обработано папок: {len(folders)}\n"
                                    f"Распаковано архивов: {total_archives}\n"
                                    f"Оригинальные архивы: {delete_status}")
            else:
                messagebox.showinfo("Готово", "Архивы не найдены или распаковать "
                                              "не удалось.\nПодробности в логе.")
        except Exception as e:
            self.log_to_gui(f"Ошибка распаковки: {e}")
            messagebox.showerror("Ошибка", f"Ошибка распаковки:\n{e}")

    # -----------------------
    # Полная обработка
    # -----------------------
    def start_processing(self):
        folders = self.get_folders()
        if not folders:
            return
        self.log_to_gui("=" * 60)
        self.log_to_gui(f"НАЧАЛО ПОЛНОЙ ОБРАБОТКИ ({len(folders)} папок)")
        self.log_to_gui("=" * 60)
        import threading
        threading.Thread(target=self._process_all, args=(folders,), daemon=True).start()

    def _process_all(self, folders):
        excel = None
        try:
            pythoncom.CoInitialize()

            # --- Этап 1: распаковка во всех папках ---
            self.set_status("Распаковка архивов...")
            self.set_progress(5)
            self.log_to_gui("Этап 1: Распаковка архивов")

            for folder_path in folders:
                self.log_to_gui(f"--- Папка: {folder_path}")
                processed = process_archives_in_folder(folder_path,
                                                       self.delete_archives_var.get())
                if processed:
                    self.log_to_gui(f"Распаковано архивов: {len(processed)}")
                    deleted_count = clean_empty_folders(folder_path)
                    if deleted_count > 0:
                        self.log_to_gui(f"Удалено {deleted_count} пустых папок")
                else:
                    self.log_to_gui("Архивы не найдены")

            # --- Этап 2: собираем файлы из всех папок ---
            self.set_status("Поиск Excel файлов...")
            self.set_progress(15)
            self.log_to_gui("Этап 2: Обработка Excel файлов")

            all_files = []
            seen = set()
            for folder_path in folders:
                found = list(folder_path.rglob("*.xls")) + list(folder_path.rglob("*.xlsx"))
                found = [f for f in found
                         if "modified" not in str(f).lower() and not f.name.startswith('~$')]
                for f in found:
                    key = str(f.resolve()).lower()   # защита от дублей при вложенных папках
                    if key not in seen:
                        seen.add(key)
                        all_files.append(f)
                self.log_to_gui(f"   {folder_path.name}: найдено {len(found)}")

            total_files = min(len(all_files), MAX_FILES_TO_PROCESS)
            if total_files == 0:
                self.set_status("Excel файлы не найдены")
                self.log_to_gui("Excel файлы не найдены")
                messagebox.showinfo("Завершено", "В выбранных папках не найдено "
                                                 "файлов .xls или .xlsx")
                return

            excel = create_excel_instance(visible=False)
            self.log_to_gui(f"Всего файлов к обработке: {total_files}")

            success_count = 0
            error_count = 0
            error_log = []

            for i, file in enumerate(all_files[:MAX_FILES_TO_PROCESS]):
                if not file.is_file():
                    continue
                self.set_status(f"Обработка {file.name} ({i + 1}/{total_files})")
                self.set_progress(15 + (i * 85 / total_files))
                self.log_to_gui(f"   Обработка: {file.name}")

                start_time = time.time()
                try:
                    result = process_excel_file(str(file), excel_app=excel)
                    elapsed = time.time() - start_time
                    if result:
                        success_count += 1
                        self.log_to_gui(f"   Успешно ({elapsed:.1f} сек)")
                    else:
                        error_count += 1
                        error_log.append(str(file))
                        self.log_to_gui(f"   Ошибка ({elapsed:.1f} сек)")
                except Exception as e:
                    error_count += 1
                    error_log.append(f"{file} - {e}")
                    self.log_to_gui(f"   Критическая ошибка: {e}")

            # лог ошибок кладём в первую из выбранных папок
            if error_log:
                log_dir = folders[0] / "modified" / "logs"
                log_dir.mkdir(parents=True, exist_ok=True)
                log_file = log_dir / "error_log.txt"
                with open(log_file, 'w', encoding='utf-8') as f:
                    f.write("Файлы, при обработке которых произошли ошибки:\n")
                    f.write("\n".join(error_log))
                self.log_to_gui(f"Лог ошибок сохранен: {log_file}")

            self.set_status("Готово")
            self.set_progress(100)
            self.log_to_gui("=" * 60)
            self.log_to_gui("ОБРАБОТКА ЗАВЕРШЕНА")
            self.log_to_gui(f"Папок: {len(folders)}")
            self.log_to_gui(f"Успешно: {success_count}")
            self.log_to_gui(f"Ошибки: {error_count}")
            self.log_to_gui("=" * 60)
            messagebox.showinfo("Завершено",
                                f"Обработка завершена!\n\n"
                                f"Папок: {len(folders)}\n"
                                f"Успешно: {success_count}\n"
                                f"Ошибки: {error_count}\n\n"
                                f"Результаты — в подпапках 'modified'")

        except Exception as ex:
            self.log_to_gui(f"Критическая ошибка: {ex}")
            messagebox.showerror("Ошибка", f"Процесс прерван:\n{ex}")
            traceback.print_exc()
        finally:
            if excel:
                try:
                    excel.Quit()
                except Exception:
                    pass
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass

    def on_close(self):
        if messagebox.askokcancel("Выход", "Вы действительно хотите выйти?"):
            self.root.destroy()

    def run(self):
        self.root.mainloop()


# -----------------------
# Запуск
# -----------------------
if __name__ == "__main__":
    try:
        logger.info("=== Запуск скрипта обработки Excel ===")
        find_archive_tool()
        app = ExcelProcessor()
        app.run()
    except KeyboardInterrupt:
        logger.info("Остановка по инициативе пользователя.")
    except Exception as e:
        logger.error(f"Произошла непредвиденная ошибка: {e}")
        traceback.print_exc()