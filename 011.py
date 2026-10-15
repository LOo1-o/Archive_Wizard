import logging
import time
import tkinter as tk
import traceback
import zipfile
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

import pythoncom
import win32com.client as win32
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

try:
    import rarfile
except ImportError:
    rarfile = None
# Настройка логгера
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger(__name__)

PAGE_WIDTH_A4_PORTRAIT = 6.5
PAGE_HEIGHT_A4 = 11.69

EXCEL_POINT_TO_INCH = 1 / 72
CALIBRATION_FACTOR = 1.02

MAX_FILES_TO_PROCESS = 1000
def process_excel_file(file_path, log_func=None):
    """Обработка одного Excel-файла, включая автоматический выбор ориентации и масштабирование."""

    def log(msg):
        if log_func:
            log_func(msg)
        else:
            print(msg)

    temp_xlsx = None
    excel = None
    wb = None

    try:
        file_path = str(file_path)
        original_path_for_modified = Path(file_path)
        if not os.path.exists(file_path):
            log(f"Файл не существует: {file_path}")
            return False

        # Конвертация .xls → .xlsx
        if file_path.lower().endswith('.xls'):
            temp_xlsx = file_path + 'x'
            try:
                log(f"Конвертируем {file_path} -> {temp_xlsx}")
                excel = create_excel_instance()
                excel.Visible = False
                excel.DisplayAlerts = False
                wb_com = excel.Workbooks.Open(file_path)
                wb_com.SaveAs(temp_xlsx, FileFormat=51)
                wb_com.Close()
                excel.Quit()
                excel = None
                file_path_for_openpyxl = temp_xlsx
            except Exception as es:
                log(f"Ошибка конвертации: {es}")
                if excel:
                    try:
                        excel.Quit()
                    except:
                        pass
                if temp_xlsx and os.path.exists(temp_xlsx):
                    try:
                        os.remove(temp_xlsx)
                    except:
                        pass
                return False
        else:
            file_path_for_openpyxl = file_path

        # Загрузка книги
        try:
            wb = load_workbook(filename=file_path_for_openpyxl, read_only=False, keep_vba=False)
        except Exception as ec:
            log(f"Ошибка загрузки: {ec}")
            return False

        # Обработка каждого листа
        for ws in wb.worksheets:
            try:
                if ws.max_row == 1 and ws.max_column == 1:
                    continue

                ws.page_setup.paperSize = ws.PAPERSIZE_A4
                # Настройка полей страницы
                ws.page_margins = PageMargins(
                    left=0.25, right=0.25,
                    top=0.5, bottom=0.25,
                    header=0.25, footer=0.25
                )

                # Рассчитываем доступную область печати
                printable_width, printable_height, orientation = calculate_printable_area(ws)

                # Рассчитываем физические размеры таблицы
                total_width, total_height = calculate_table_dimensions(ws)
                log(f"[{ws.title}] Физические размеры таблицы: {total_width:.2f} x {total_height:.2f} дюймов")
                log(
                    f"[{ws.title}] Доступная область печати: {printable_width:.2f} x {printable_height:.2f} дюймов ({orientation})")

                # Оптимизация столбцов перед расчетами
                auto_adjust_columns(ws, max_width=40)

                # Пересчитываем размеры после оптимизации
                total_width, total_height = calculate_table_dimensions(ws)
                log(f"[{ws.title}] Размеры после оптимизации: {total_width:.2f} x {total_height:.2f} дюймов")

                # Определяем ориентацию
                if orientation == "portrait":
                    available_width = printable_width
                    available_height = printable_height
                else:
                    available_width = printable_width
                    available_height = printable_height

                # Автоматический выбор ориентации
                if total_width <= available_width:
                    # Таблица помещается в текущую ориентацию
                    ws.page_setup.fitToWidth = 1
                    ws.page_setup.fitToHeight = 0
                    ws.page_setup.scale = None
                    log(f"[{ws.title}] Таблица помещается ({total_width:.2f} <= {available_width:.2f} дюймов)")
                else:
                    # Пробуем изменить ориентацию
                    if orientation == "portrait":
                        ws.page_setup.orientation = "landscape"
                        # Пересчитываем доступную область после смены ориентации
                        printable_width, printable_height, orientation = calculate_printable_area(ws)
                        available_width = printable_width
                        log(f"[{ws.title}] Переключено на альбомную ориентацию")

                    if total_width <= available_width:
                        ws.page_setup.fitToWidth = 1
                        ws.page_setup.fitToHeight = 0
                        ws.page_setup.scale = None
                        log(
                            f"[{ws.title}] Таблица помещается после смены ориентации ({total_width:.2f} <= {available_width:.2f} дюймов)")
                    else:
                        # Рассчитываем необходимый масштаб
                        required_scale = (available_width / total_width) * 100
                        final_scale = max(10, min(100, int(required_scale)))

                        # Применяем явное масштабирование
                        ws.page_setup.fitToWidth = 0
                        ws.page_setup.fitToHeight = 0
                        ws.page_setup.scale = final_scale

                        log(f"[{ws.title}] Установлен масштаб: {final_scale}% (требовалось: {required_scale:.1f}%)")

                        # Предупреждение для экстремальных случаев
                        if final_scale < 50:
                            log(f"ВНИМАНИЕ! Таблица слишком широкая ({ws.max_column} столбцов)")
                            log("Рекомендации:")
                            log("- Проверьте наличие пустых столбцов")
                            log("- Рассмотрите возможность разделения таблицы")
                            log("- Уменьшите поля страницы")

                # Проверяем высоту таблицы
                if total_height > available_height:
                    log(
                        f"[{ws.title}] ВНИМАНИЕ! Таблица не помещается по высоте ({total_height:.2f} > {available_height:.2f} дюймов)")
                    log("Рекомендуется уменьшить количество строк или размер шрифта")

                log(f"[{ws.title}] Окончательные настройки:")
                log(f"Ориентация: {'альбомная' if ws.page_setup.orientation == 'landscape' else 'портретная'}")
                log(f"fitToWidth: {ws.page_setup.fitToWidth}")
                log(f"Масштаб: {ws.page_setup.scale or 'авто'}")

            except Exception as ev:
                log(f"Ошибка обработки листа {ws.title}: {str(ev)}")
                continue

        # Сохраняем в папку modified
        try:
            modified_path = get_modified_path(original_path_for_modified,
                                              os.path.dirname(original_path_for_modified)).with_suffix('.xlsx')
            modified_path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(str(modified_path))
            log(f"Сохранено как .xlsx: {modified_path}")
            return True
        except Exception as ez:
            log(f"Ошибка сохранения: {ez}")
            return False

    finally:
        if wb:
            try:
                wb.close()
            except Exception:
                pass
        if temp_xlsx and os.path.exists(temp_xlsx):
            try:
                os.remove(temp_xlsx)
                log(f"Удалён временный файл: {temp_xlsx}")
            except PermissionError as ep:
                log(f"Не удалось удалить временный файл из-за ошибки доступа: {ep}")
            except FileNotFoundError:
                pass
            except OSError as es:
                log(f"Не удалось удалить временный файл: {es}")


def auto_adjust_columns(ws, max_width=30):
    """
    Автоматически оптимизирует ширину столбцов
    :param ws: Рабочий лист
    :param max_width: Максимальная ширина столбца в символах
    """
    for col in ws.columns:
        if not col:
            continue

        column_letter = col[0].column_letter
        col_dim = ws.column_dimensions[column_letter]

        # Пропускаем скрытые столбцы
        if col_dim.hidden:
            continue

        # Ищем максимальную длину содержимого
        max_length = 0
        for cell in col:
            try:
                if cell.value:
                    # Учитываем многострочные значения
                    lines = str(cell.value).split('\n')
                    for line in lines:
                        if len(line) > max_length:
                            max_length = len(line)
            except:
                pass

        # Устанавливаем ширину с ограничением
        if max_length > 0:
            adjusted_width = min(max_length + 2, max_width)  # +2 для отступов
            col_dim.width = adjusted_width


def calculate_printable_area(ws):
    """Рассчитывает доступную область печати в дюймах с учетом ориентации"""
    # 1. Получаем настройки страницы
    page_setup = ws.page_setup
    margins = ws.page_margins

    # 2. Определяем размер бумаги (в дюймах)
    paper_sizes = {
        1: (8.5, 11),  # Letter
        9: (8.27, 11.69),  # A4
        5: (8.5, 14),  # Legal
        # Добавьте другие форматы при необходимости
    }

    paper_code = page_setup.paperSize or 9  # По умолчанию A4
    paper_width, paper_height = paper_sizes.get(paper_code, (8.27, 11.69))

    # 3. Учитываем ориентацию
    orientation = "portrait"
    if page_setup.orientation == "landscape":
        paper_width, paper_height = paper_height, paper_width
        orientation = "landscape"

    # 4. Рассчитываем доступную область
    printable_width = paper_width - (margins.left + margins.right)
    printable_height = paper_height - (margins.top + margins.bottom)

    return printable_width, printable_height, orientation