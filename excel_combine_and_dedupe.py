import openpyxl
import os
import win32com.client as win32
from tkinter import filedialog, Tk
from glob import glob

def copy_cell_format(source_cell, target_cell):
    """Копирует форматирование из source_cell в target_cell"""
    if source_cell.has_style:
        target_cell.font = source_cell.font.copy()
        target_cell.border = source_cell.border.copy()
        target_cell.fill = source_cell.fill.copy()
        target_cell.number_format = source_cell.number_format
        target_cell.protection = source_cell.protection.copy()
        target_cell.alignment = source_cell.alignment.copy()

# ---- 1. Выбор папки ----
root = Tk()
root.withdraw()
folder_path = filedialog.askdirectory(title="Выберите папку с Excel-файлами")
if not folder_path:
    print("Папка не выбрана.")
    exit()

files = glob(os.path.join(folder_path, "*.xlsx"))
if not files:
    print("Нет xlsx-файлов.")
    exit()

print(f"Найдено файлов: {len(files)}")

# ---- 2. Объединение с сохранением форматирования ----
target_wb = openpyxl.Workbook()
target_ws = target_wb.active
target_ws.title = "Объединённые данные"

first_file = True
target_row = 1

for file in files:
    print(f"Обработка: {os.path.basename(file)}")
    wb = openpyxl.load_workbook(file)
    ws = wb.active

    if first_file:
        # Копируем заголовки (первая строка) из первого файла
        for col_idx, cell in enumerate(ws[1], start=1):
            target_cell = target_ws.cell(row=target_row, column=col_idx, value=cell.value)
            copy_cell_format(cell, target_cell)
        target_row += 1
        first_file = False
        start_row = 2
    else:
        start_row = 2  # пропускаем заголовок

    # Копируем данные с форматированием
    for row in ws.iter_rows(min_row=start_row):
        for col_idx, cell in enumerate(row, start=1):
            target_cell = target_ws.cell(row=target_row, column=col_idx, value=cell.value)
            copy_cell_format(cell, target_cell)
        target_row += 1

    wb.close()

output_file = os.path.join(folder_path, "merged_with_format.xlsx")
target_wb.save(output_file)
print(f"Файл с форматированием сохранён: {output_file}")

# ---- 3. Автоматическое удаление дубликатов через Excel ----
print("Запуск Excel для удаления дубликатов...")
excel = win32.gencache.EnsureDispatch('Excel.Application')
excel.Visible = False  # False - Excel не будет показан, True - можно наблюдать процесс

try:
    wb = excel.Workbooks.Open(output_file)
    ws = wb.ActiveSheet

    # Определяем последнюю заполненную строку в столбце A (ОКПО)
    last_row = ws.Cells(ws.Rows.Count, 1).End(-4162).Row  # -4162 = xlUp
    # Диапазон от A1 до C(last_row)
    range_address = f"A1:C{last_row}"
    rng = ws.Range(range_address)

    # Удаляем дубликаты по первому столбцу (ОКПО), с заголовками
    rng.RemoveDuplicates(Columns=1, Header=1)  # Header=1 означает xlYes

    wb.Save()
    wb.Close()
    print("Дубликаты успешно удалены.")
except Exception as e:
    print(f"Ошибка при работе с Excel: {e}")
finally:
    excel.Quit()

print("Готово! Итоговый файл:", output_file)