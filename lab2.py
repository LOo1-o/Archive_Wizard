#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import logging
import os
import time
import tkinter as tk
import traceback
from tkinter import filedialog, messagebox, ttk
from pathlib import Path
import pythoncom
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import re
import win32com.client as win32
from openpyxl.worksheet.page import PageMargins

# Настройка логгера
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Константы
PAGE_WIDTH_A4_PORTRAIT = 8.27
PAGE_HEIGHT_A4_PORTRAIT = 11.69
EXCEL_CHAR_TO_INCH = 1 / 13
EXCEL_POINT_TO_INCH = 1 / 72
CHAR_WIDTH_MULTIPLIER = 1.23
MAX_FILES_TO_PROCESS = 1000


# ===================== ФУНКЦИИ ДЛЯ РАБОТЫ С ЗАГОЛОВКАМИ =====================
def find_header_rows(ws, first_data_row=1):
    """Находит строки заголовков таблицы"""
    try:
        marker_row = None
        for col in [1, 2]:
            for row in range(first_data_row, min(first_data_row + 20, ws.max_row + 1)):
                cell = ws.cell(row=row, column=col)
                cell_str = str(cell.value).strip().lower() if cell.value else ""

                if cell_str in ['a', 'b', 'c', 'а', 'б', 'п/п', '№'] or re.fullmatch(r'^[123]$', cell_str):
                    marker_row = row
                    break
            if marker_row:
                break

        if not marker_row:
            marker_row = first_data_row
            logger.info(f"Маркер не найден, используется строка: {marker_row}")
            return marker_row, marker_row

        header_block = [marker_row]
        current_row = marker_row - 1

        while current_row >= 1:
            if any(ws.cell(row=current_row, column=col).border.right.style for col in
                   range(1, min(10, ws.max_column + 1))):
                header_block.append(current_row)
                current_row -= 1
            else:
                break

        start_row, end_row = min(header_block), max(header_block)
        logger.info(f"Сквозные строки: {start_row}-{end_row} + Сквозной столбец A")
        return start_row, end_row

    except Exception as e:
        logger.error(f"Ошибка поиска заголовков: {e}")
        return 1, 1


def safe_add_page_number(ws):
    """Добавляет номер страницы в верхний колонтитул"""
    try:
        if not (ws.oddHeader.center and "&P" in getattr(ws.oddHeader.center, 'text', '')):
            ws.oddHeader.center.text = ws.evenHeader.center.text = "Страница &P из &N"
            return True
        return False
    except Exception as e:
        logger.error(f"Ошибка добавления номера страницы: {e}")
        return False


# ===================== ФУНКЦИИ ДЛЯ НАСТРОЙКИ ПЕЧАТИ =====================
def calculate_table_dimensions(ws):
    """
    Рассчитывает размеры таблицы в дюймах.
    Возвращает кортеж (ширина, высота).
    """
    # Определяем последний непустой столбец
    max_col = max(
        (cell.column for row in ws.iter_rows() for cell in row if cell.value),
        default=1
    )

    # Рассчитываем ширину
    total_width = sum(
        (col.width * EXCEL_CHAR_TO_INCH
         if (col := ws.column_dimensions.get(get_column_letter(col_idx))) and col.width is not None
         else 8.43 * EXCEL_CHAR_TO_INCH)
        for col_idx in range(1, max_col + 1)
        if not ws.column_dimensions.get(get_column_letter(col_idx), {}).hidden
    )

    # Рассчитываем высоту (ИСКЛЮЧАЕМ скрытые строки полностью)
    total_height = sum(
        (row.height * EXCEL_POINT_TO_INCH
         if (row := ws.row_dimensions.get(row_idx)) and row.height is not None
         else 15.0 * EXCEL_POINT_TO_INCH)
        for row_idx in range(1, ws.max_row + 1)
        if not ws.row_dimensions.get(row_idx, {}).hidden
    )

    return total_width, total_height


def calculate_printable_area(ws):
    """Рассчитывает доступную область печати"""
    paper_sizes = {1: (8.5, 11), 9: (8.27, 11.69), 5: (8.5, 14)}
    paper_code = ws.page_setup.paperSize or 9
    paper_width, paper_height = paper_sizes.get(paper_code, (8.27, 11.69))

    orientation = "landscape" if ws.page_setup.orientation == "landscape" else "portrait"
    if orientation == "landscape":
        paper_width, paper_height = paper_height, paper_width

    margins = ws.page_margins
    printable_width = paper_width - (margins.left + margins.right)
    printable_height = paper_height - (margins.top + margins.bottom)

    return printable_width, printable_height, orientation


"""def auto_adjust_columns(ws, max_width=30):
   Оптимизирует ширину столбцов без учета принудительных переносов текста
    for col in ws.columns:
        if not col or ws.column_dimensions[col[0].column_letter].hidden:
            continue

        # Ищем максимальную длину значения в столбце (без разбивки на строки)
        max_length = max(
            len(str(cell.value)) if cell.value else 0
            for cell in col
        )

        if max_length > 0:
            # Устанавливаем ширину с небольшим запасом (+2 символа)
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, max_width)"""


def adjust_sheet_print_settings(ws, log_func=None):
    """Настраивает ориентацию и масштабирование для рабочего листа"""

    def log(msg):
        if log_func:
            log_func(msg)
        else:
            print(msg)

    try:
        if ws.max_row == 1 and ws.max_column == 1:
            log(f"[{ws.title}] Лист пустой, пропускаем")
            return

        # Устанавливаем стандартные настройки страницы
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_margins = PageMargins(
            left=0.25, right=0.25,
            top=0.5, bottom=0.25,
            header=0.25, footer=0.25
        )

        # Оптимизация столбцов перед расчетами
        #auto_adjust_columns(ws, max_width=40)

        # Рассчитываем доступную область печати
        printable_width, printable_height, orientation = calculate_printable_area(ws)

        # Рассчитываем физические размеры таблицы
        total_width, total_height = calculate_table_dimensions(ws)
        log(f"[{ws.title}] Физические размеры таблицы: {total_width:.2f} x {total_height:.2f} дюймов")
        log(
            f"[{ws.title}] Доступная область печати: {printable_width:.2f} x {printable_height:.2f} дюймов ({orientation})")

        # Определяем ориентацию
        if orientation == "portrait":
            available_width = printable_width
            available_height = printable_height
        else:
            available_width = printable_width
            available_height = printable_height

        # Автоматический выбор ориентации
        if total_width <= available_width:
            # Таблица помещается в текущую ориентацию
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.page_setup.scale = None
            log(f"[{ws.title}] Таблица помещается ({total_width:.2f} <= {available_width:.2f} дюймов)")
        else:
            # Пробуем изменить ориентацию
            if orientation == "portrait":
                ws.page_setup.orientation = "landscape"
                # Пересчитываем доступную область после смены ориентации
                printable_width, printable_height, orientation = calculate_printable_area(ws)
                available_width = printable_width
                log(f"[{ws.title}] Переключено на альбомную ориентацию")

            if total_width <= available_width:
                ws.page_setup.fitToWidth = 1
                ws.page_setup.fitToHeight = 0
                ws.page_setup.scale = None
                log(
                    f"[{ws.title}] Таблица помещается после смены ориентации ({total_width:.2f} <= {available_width:.2f} дюймов)")
            else:
                # Рассчитываем необходимый масштаб
                required_scale = (available_width / total_width) * 100
                final_scale = max(60, min(100, int(required_scale)))

                # Применяем явное масштабирование
                ws.page_setup.fitToWidth = 0
                ws.page_setup.fitToHeight = 0
                ws.page_setup.scale = final_scale

                log(f"[{ws.title}] Установлен масштаб: {final_scale}% (требовалось: {required_scale:.1f}%)")

                # Предупреждение для экстремальных случаев
                if final_scale < 50:
                    log(f"ВНИМАНИЕ! Таблица слишком широкая ({ws.max_column} столбцов)")
                    log("Рекомендации:")
                    log("- Проверьте наличие пустых столбцов")
                    log("- Рассмотрите возможность разделения таблицы")
                    log("- Уменьшите поля страницы")

        # Проверяем высоту таблицы
        if total_height > available_height:
            log(
                f"[{ws.title}] ВНИМАНИЕ! Таблица не помещается по высоте ({total_height:.2f} > {available_height:.2f} дюймов)")
            log("Рекомендуется уменьшить количество строк или размер шрифта")

        log(f"[{ws.title}] Окончательные настройки:")
        log(f"Ориентация: {'альбомная' if ws.page_setup.orientation == 'landscape' else 'портретная'}")
        log(f"fitToWidth: {ws.page_setup.fitToWidth}")
        log(f"Масштаб: {ws.page_setup.scale or 'авто'}")

    except Exception as ev:
        log(f"Ошибка обработки листа {ws.title}: {str(ev)}")
        raise  # Пробрасываем исключение для обработки в вызывающем коде


# ===================== ФУНКЦИИ ДЛЯ КОНВЕРТАЦИИ ФАЙЛОВ =====================
def create_excel_instance():
    """Создает экземпляр Excel"""
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    return excel


def convert_xls_to_xlsx(file_path, log_func=None):
    """Конвертирует XLS в XLSX"""

    def log(msg):
        log_func(msg) if log_func else print(msg)

    temp_xlsx = file_path + 'x'
    excel = None

    try:
        log(f"Конвертация {file_path} -> {temp_xlsx}")
        excel = create_excel_instance()
        excel.DisplayAlerts = False
        wb = excel.Workbooks.Open(file_path)
        wb.SaveAs(temp_xlsx, FileFormat=51)
        return temp_xlsx
    except Exception as e:
        log(f"Ошибка конвертации: {e}")
        raise
    finally:
        if excel:
            try:
                excel.Quit()
            except:
                pass


# ===================== ФУНКЦИИ ДЛЯ РАБОТЫ С WIN32 =====================
def setup_print_titles_win32(file_path, header_rows_map, log_func=None):
    """Настраивает параметры печати через Win32"""

    def log(msg):
        log_func(msg) if log_func else print(msg)

    excel = None
    try:
        excel = create_excel_instance()
        wb = excel.Workbooks.Open(file_path)

        for ws in wb.Sheets:
            try:
                if ws.Name in header_rows_map:
                    start, end = header_rows_map[ws.Name]
                    ws.PageSetup.PrintTitleRows = f"${start}:${end}"
                    ws.PageSetup.PrintTitleColumns = "$A:$A"
                    log(f"[{ws.Name}] Win32: сквозные строки {start}-{end}")

                ws.PageSetup.Order = 2
                ws.PageSetup.FitToPagesWide = 1
                ws.PageSetup.FitToPagesTall = False
                log(f"[{ws.Name}] Win32: порядок печати установлен")

            except Exception as e:
                log(f"[{ws.Name}] Win32 ошибка: {e}")

        wb.Save()
        return True
    except Exception as e:
        log(f"Win32COM ошибка: {e}")
        return False
    finally:
        if excel:
            try:
                excel.Quit()
            except:
                pass


# ===================== ОСНОВНАЯ ФУНКЦИЯ ОБРАБОТКИ =====================
def process_excel_file(file_path, log_func=None):
    """Обрабатывает Excel-файл"""

    def log(msg):
        log_func(msg) if log_func else print(msg)

    temp_xlsx = None
    wb = None
    header_rows_map = {}

    try:
        file_path = Path(file_path)
        if not file_path.exists():
            log(f"Файл не существует: {file_path}")
            return False

        # Конвертация XLS -> XLSX
        if file_path.suffix.lower() == '.xls':
            try:
                temp_xlsx = convert_xls_to_xlsx(str(file_path), log_func=log)
                file_to_process = temp_xlsx
            except Exception:
                return False
        else:
            file_to_process = str(file_path)

        # Обработка XLSX
        try:
            wb = load_workbook(filename=file_to_process, read_only=False, keep_vba=False)
        except Exception as e:
            log(f"Ошибка загрузки: {e}")
            return False

        # Обработка листов
        for ws in wb.worksheets:
            try:
                if ws.max_row == 1 and ws.max_column == 1:
                    continue

                # Настройка заголовков
                start_row, end_row = find_header_rows(ws)
                header_rows_map[ws.title] = (start_row, end_row)
                ws.print_title_rows = f"{start_row}:{end_row}"
                ws.print_title_cols = "A:A"
                log(f"[{ws.title}] Сквозные строки: {start_row}-{end_row}")

                # Настройка печати
                safe_add_page_number(ws)
                adjust_sheet_print_settings(ws, log_func=log)

            except Exception as e:
                log(f"[{ws.title}] Ошибка обработки: {e}")

        # Сохранение
        modified_path = get_modified_path(file_path, file_path.parent).with_suffix('.xlsx')
        modified_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(modified_path))
        log(f"Сохранено: {modified_path}")

        # Win32 настройки
        setup_print_titles_win32(str(modified_path), header_rows_map, log_func=log)
        return True

    finally:
        if wb:
            try:
                wb.close()
            except:
                pass
        if temp_xlsx and os.path.exists(temp_xlsx):
            try:
                os.remove(temp_xlsx)
                log(f"Удален временный файл: {temp_xlsx}")
            except Exception as e:
                log(f"Ошибка удаления временного файла: {e}")


# ===================== GUI И ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ =====================
class ExcelProcessor:
    """GUI для обработки файлов"""

    def __init__(self):
        self.root = tk.Tk()
        self.root.withdraw()
        self.folder_path = tk.StringVar()
        self.progress_var = tk.DoubleVar()
        self.status_var = tk.StringVar(value="Готово")
        self.log_messages = []
        self.setup_ui()

    def setup_ui(self):
        self.root.deiconify()
        self.root.title("Wizard Archive")
        self.root.geometry("600x400")

        main_frame = ttk.Frame(self.root, padding="20")
        main_frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(main_frame, text="Wizard Archive", font=('Helvetica', 14, 'bold')).pack(pady=10)
        ttk.Label(main_frame, text="Папка с файлами:").pack(anchor="w")
        ttk.Entry(main_frame, textvariable=self.folder_path, width=50).pack(fill="x", pady=5)
        ttk.Button(main_frame, text="Обзор", command=self.select_folder).pack(pady=5)
        ttk.Button(main_frame, text="Старт", command=self.start_processing).pack(pady=20)

        self.progress_bar = ttk.Progressbar(main_frame, variable=self.progress_var, maximum=100)
        self.progress_bar.pack(fill="x")
        ttk.Label(main_frame, textvariable=self.status_var).pack()

        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    def log_message(self, message):
        print(message)
        self.log_messages.append(message)

    def select_folder(self):
        folder = filedialog.askdirectory()
        if folder:
            self.folder_path.set(folder)

    def start_processing(self):
        folder = self.folder_path.get()
        if not folder:
            messagebox.showerror("Ошибка", "Сначала выберите папку!")
            return
        if not os.path.exists(folder):
            messagebox.showerror("Ошибка", "Папка не найдена!")
            return

        import threading
        threading.Thread(target=self.process_files_thread, args=(folder,), daemon=True).start()

    def process_files_thread(self, root_folder):
        try:
            pythoncom.CoInitialize()
            self.status_var.set("Подготовка...")
            self.root.update()

            success_count = 0
            error_count = 0
            error_log = []

            files = list(Path(root_folder).rglob("*.xls")) + list(Path(root_folder).rglob("*.xlsx"))
            files = [f for f in files if "modified" not in str(f).lower() and not f.name.startswith('~$')]
            total_files = min(len(files), MAX_FILES_TO_PROCESS)

            if total_files == 0:
                self.status_var.set("Файлы для обработки не найдены.")
                messagebox.showinfo("Завершено", "В выбранной папке и ее подпапка не найдено файлов .xls или .xlsx.")
                return

            for i, file in enumerate(files):
                if not file.is_file():
                    continue

                self.status_var.set(f"Обработка {file.name}...")
                self.progress_var.set(((i + 1) / total_files) * 100)
                self.root.update()

                start_time = time.time()
                result = process_excel_file(str(file), log_func=lambda msg: print(f"[{file.name}] {msg}"))
                elapsed = time.time() - start_time

                if result:
                    success_count += 1
                    print(f"[{file.name}] ✅ Обработан за {elapsed:.2f} секунд")
                else:
                    error_count += 1
                    error_log.append(str(file))
                    print(f"[{file.name}] ❌ Ошибка за {elapsed:.2f} секунд")

            if error_log:
                log_dir = Path(root_folder) / "modified" / "logs"
                log_dir.mkdir(parents=True, exist_ok=True)
                log_file = log_dir / "error_log.txt"
                with open(log_file, 'w', encoding='utf-8') as f:
                    f.write("Файлы, при обработке которых произошли ошибки:\n")
                    f.write("\n".join(error_log))

            messagebox.showinfo("Завершено", f"Успешно: {success_count}, Ошибки: {error_count}")

        except Exception as ex:
            messagebox.showerror("Критическая ошибка", f"Процесс прерван:\n{str(ex)}")
            traceback.print_exc()
        finally:
            self.status_var.set("Готово")
            self.progress_var.set(0)
            try:
                pythoncom.CoUninitialize()
            except:
                pass

    def on_close(self):
        if messagebox.askokcancel("Выход", "Вы действительно хотите выйти?"):
            self.root.destroy()

    def run(self):
        self.root.mainloop()


def get_modified_path(original_path, root_folder):
    """Генерирует путь для сохранения модифицированных файлов"""
    original_path = Path(original_path)
    try:
        rel_path = original_path.relative_to(
            root_folder) if root_folder in original_path.parents else original_path.name
    except ValueError:
        rel_path = original_path.name
    return original_path.parent / "modified" / rel_path


if __name__ == "__main__":
    try:
        app = ExcelProcessor()
        app.run()
    except KeyboardInterrupt:
        print("Остановка по инициативе пользователя.")
    except Exception as e:
        print(f"Произошла непредвиденная ошибка: {e}")
        traceback.print_exc()